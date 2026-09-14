from __future__ import annotations

import logging
from dataclasses import replace
from io import BytesIO

import pytest
from openpyxl import Workbook

from gdrive_scoped.documents import DocumentService
from gdrive_scoped.drive import FOLDER_MIME_TYPE, DriveItem, SearchPage
from gdrive_scoped.errors import EmptyDocument, ExportTooLarge, InvalidCursor, ScopeViolation
from gdrive_scoped.extractors import ExtractedDocument, ExtractorRegistry, TextBlock
from gdrive_scoped.extractors.registry import ExtractionPlan
from gdrive_scoped.location import DriveKind, DriveLocation
from gdrive_scoped.scope import ScopedDrive, audit_caller


def _registry_returning(document: ExtractedDocument) -> ExtractorRegistry:
    """A registry whose PDF plan yields exactly this parse."""

    registry = ExtractorRegistry()
    registry._plans["application/pdf"] = ExtractionPlan(  # noqa: SLF001
        export_mime_type=None, extractor=lambda _: document
    )
    return registry


SHARED_DRIVE = DriveLocation(DriveKind.SHARED_DRIVE, "drive")
MY_DRIVE = DriveLocation(DriveKind.MY_DRIVE)


class SearchGateway:
    def __init__(
        self,
        *items: DriveItem,
        search_results: list[DriveItem],
        blobs: dict[str, bytes] | None = None,
        has_more: bool = False,
    ) -> None:
        self.items = {item.id: item for item in items}
        self.search_results = search_results
        self.has_more = has_more
        self.searched_limits: list[int] = []
        self.blobs = blobs or {}
        self.searched_parent_batches: list[tuple[str, ...]] = []
        self.folder_enumerations = 0
        self.child_listings: list[str] = []
        self.metadata_reads: list[str] = []
        self.exports: list[tuple[str, str]] = []
        self.downloads: list[str] = []

    async def get_item(self, item_id: str) -> DriveItem:
        self.metadata_reads.append(item_id)
        return self.items[item_id]

    async def list_children(self, folder_id: str) -> list[DriveItem]:
        self.child_listings.append(folder_id)
        return [item for item in self.items.values() if item.parents == (folder_id,)]

    async def list_folders(self) -> list[DriveItem]:
        self.folder_enumerations += 1
        return [item for item in self.items.values() if item.mime_type == FOLDER_MIME_TYPE]

    async def list_descendants(self, parent_ids: tuple[str, ...]) -> list[DriveItem]:
        # The same query without the keyword clause, which is what it is.
        return (await self.search_items(parent_ids, "")).items

    async def search_items(
        self, parent_ids: tuple[str, ...], keywords: str, limit: int = 1000
    ) -> SearchPage:
        self.searched_parent_batches.append(parent_ids)
        self.searched_limits.append(limit)
        return SearchPage(items=self.search_results, has_more=self.has_more)

    async def download_item(self, item_id: str) -> bytes:
        self.downloads.append(item_id)
        return self.blobs[item_id]

    async def export_item(self, item_id: str, mime_type: str) -> bytes:
        self.exports.append((item_id, mime_type))
        return self.blobs[item_id]


class BatchSearchGateway(SearchGateway):
    def __init__(
        self,
        *items: DriveItem,
        first_batch: list[DriveItem],
        second_batch: list[DriveItem],
    ) -> None:
        super().__init__(*items, search_results=[])
        self._batch_results = [first_batch, second_batch]

    async def list_descendants(self, parent_ids: tuple[str, ...]) -> list[DriveItem]:
        # The same query without the keyword clause, which is what it is.
        return (await self.search_items(parent_ids, "")).items

    async def search_items(
        self, parent_ids: tuple[str, ...], keywords: str, limit: int = 1000
    ) -> SearchPage:
        self.searched_parent_batches.append(parent_ids)
        return SearchPage(items=self._batch_results[len(self.searched_parent_batches) - 1])


@pytest.mark.anyio
async def test_keyword_search_covers_nested_folders_and_returns_relative_paths() -> None:
    root = DriveItem("root-folder", "Corpus", FOLDER_MIME_TYPE, "drive")
    finance = DriveItem("finance", "Finance", FOLDER_MIME_TYPE, "drive", parents=("root-folder",))
    plans = DriveItem("plans", "Plans", FOLDER_MIME_TYPE, "drive", parents=("finance",))
    report = DriveItem("report", "Budget.md", "text/markdown", "drive", parents=("plans",))
    gateway = SearchGateway(root, finance, plans, report, search_results=[report, report])
    service = DocumentService(ScopedDrive(gateway, SHARED_DRIVE, "root-folder"))

    results = await service.search_documents("budget")

    assert [(result.name, result.relative_path) for result in results.items] == [
        ("Budget.md", "Finance/Plans/Budget.md")
    ]
    assert gateway.searched_parent_batches == [("root-folder", "finance", "plans")]
    assert results.items[0].id == report.id


@pytest.mark.anyio
async def test_keyword_search_in_my_drive_drops_cross_location_results() -> None:
    root = DriveItem("root-folder", "Corpus", FOLDER_MIME_TYPE, None)
    plans = DriveItem("plans", "Plans", FOLDER_MIME_TYPE, None, parents=("root-folder",))
    report = DriveItem("report", "Budget.md", "text/markdown", None, parents=("plans",))
    cross_location = DriveItem(
        "shared-report",
        "Shared budget.md",
        "text/markdown",
        "drive",
        parents=("plans",),
    )
    gateway = SearchGateway(
        root,
        plans,
        report,
        cross_location,
        search_results=[report, cross_location],
    )
    service = DocumentService(ScopedDrive(gateway, MY_DRIVE, "root-folder"))

    results = await service.search_documents("budget")

    assert [(result.name, result.relative_path) for result in results.items] == [
        ("Budget.md", "Plans/Budget.md")
    ]
    assert gateway.searched_parent_batches == [("root-folder", "plans")]


@pytest.mark.anyio
async def test_keyword_search_chunks_large_folder_sets_before_calling_drive() -> None:
    root = DriveItem("root-folder", "Corpus", FOLDER_MIME_TYPE, "drive")
    folders = [
        DriveItem(
            f"folder-{index}",
            f"Folder {index}",
            FOLDER_MIME_TYPE,
            "drive",
            parents=("root-folder",),
        )
        for index in range(401)
    ]
    gateway = SearchGateway(root, *folders, search_results=[])
    service = DocumentService(ScopedDrive(gateway, SHARED_DRIVE, "root-folder"))

    await service.search_documents("budget")

    assert [len(batch) for batch in gateway.searched_parent_batches] == [400, 2]


@pytest.mark.anyio
async def test_keyword_search_keeps_drive_relevance_order_within_one_batch() -> None:
    root = DriveItem("root-folder", "Corpus", FOLDER_MIME_TYPE, "drive")
    older = DriveItem(
        "older",
        "Older.md",
        "text/markdown",
        "drive",
        parents=("root-folder",),
        modified_time="2026-01-01T00:00:00Z",
    )
    newer = DriveItem(
        "newer",
        "Newer.md",
        "text/markdown",
        "drive",
        parents=("root-folder",),
        modified_time="2026-08-30T00:00:00Z",
    )
    # Drive answered with the older document first: it is the more relevant one,
    # and relevance is the only ranking signal a keyword search has.
    gateway = SearchGateway(root, older, newer, search_results=[older, newer])
    service = DocumentService(ScopedDrive(gateway, SHARED_DRIVE, "root-folder"))

    results = await service.search_documents("report")

    assert [result.name for result in results.items] == ["Older.md", "Newer.md"]


@pytest.mark.anyio
async def test_keyword_search_interleaves_parent_batches_by_rank() -> None:
    root = DriveItem("root-folder", "Corpus", FOLDER_MIME_TYPE, "drive")
    folders = [
        DriveItem(
            f"folder-{index}",
            f"Folder {index}",
            FOLDER_MIME_TYPE,
            "drive",
            parents=("root-folder",),
        )
        for index in range(401)
    ]

    def document(item_id: str, parent: str, modified_time: str) -> DriveItem:
        return DriveItem(
            item_id,
            f"{item_id}.md",
            "text/markdown",
            "drive",
            parents=(parent,),
            modified_time=modified_time,
        )

    # The first batch holds the older half of the corpus; recency would bury it.
    first_batch = [
        document("a1", "folder-0", "2026-01-01T00:00:00Z"),
        document("a2", "folder-0", "2026-01-02T00:00:00Z"),
    ]
    second_batch = [
        document("b1", "folder-400", "2026-08-01T00:00:00Z"),
        document("b2", "folder-400", "2026-08-02T00:00:00Z"),
    ]
    gateway = BatchSearchGateway(
        root,
        *folders,
        *first_batch,
        *second_batch,
        first_batch=first_batch,
        second_batch=second_batch,
    )
    service = DocumentService(ScopedDrive(gateway, SHARED_DRIVE, "root-folder"))

    results = await service.search_documents("report", limit=3)

    # Rank position is the only thing comparable across batches, so the first
    # hit of every batch comes before the second hit of any.
    assert [result.name for result in results.items] == ["a1.md", "b1.md", "a2.md"]
    assert results.truncated is True
    assert len(gateway.searched_parent_batches) == 2


@pytest.mark.anyio
@pytest.mark.parametrize(("query", "limit"), [("", 10), ("budget", 0), ("budget", 51)])
async def test_keyword_search_rejects_unbounded_or_empty_requests(query: str, limit: int) -> None:
    root = DriveItem("root-folder", "Corpus", FOLDER_MIME_TYPE, "drive")
    service = DocumentService(
        ScopedDrive(SearchGateway(root, search_results=[]), SHARED_DRIVE, "root-folder")
    )

    with pytest.raises(ValueError):
        await service.search_documents(query, limit=limit)


@pytest.mark.anyio
async def test_folder_listing_is_bounded_and_continues_with_a_cursor() -> None:
    root = DriveItem("root-folder", "Corpus", FOLDER_MIME_TYPE, "drive")
    documents = [
        DriveItem(
            f"document-{index}",
            f"Document {index}.txt",
            "text/plain",
            "drive",
            parents=("root-folder",),
        )
        for index in range(3)
    ]
    service = DocumentService(
        ScopedDrive(SearchGateway(root, *documents, search_results=[]), SHARED_DRIVE, "root-folder")
    )

    first = await service.list_folder(limit=2)
    second = await service.list_folder(cursor=first.next_cursor, limit=2)

    assert [item.name for item in first.items] == ["Document 0.txt", "Document 1.txt"]
    assert first.next_cursor is not None
    assert [item.name for item in second.items] == ["Document 2.txt"]
    assert second.next_cursor is None


@pytest.mark.anyio
async def test_plain_text_reads_continue_without_skipping_or_repeating_content() -> None:
    root = DriveItem("root-folder", "Corpus", FOLDER_MIME_TYPE, "drive")
    document = DriveItem(
        "document",
        "notes.txt",
        "text/plain",
        "drive",
        parents=("root-folder",),
        modified_time="2026-08-30T10:00:00Z",
    )
    original = b"abcdefghijklmnopqrstuvwxyz"
    gateway = SearchGateway(root, document, search_results=[], blobs={document.id: original})
    service = DocumentService(ScopedDrive(gateway, SHARED_DRIVE, "root-folder"))
    [listed] = (await service.list_folder()).items

    chunks: list[str] = []
    cursor: str | None = None
    while True:
        result = await service.read_document(listed.id, cursor=cursor, max_chars=10)
        chunks.append(result.content)
        cursor = result.next_cursor
        if cursor is None:
            break

    assert "".join(chunks) == original.decode()
    assert [len(chunk) for chunk in chunks] == [10, 10, 6]
    assert gateway.downloads == [document.id]


@pytest.mark.anyio
async def test_google_docs_are_exported_as_markdown() -> None:
    root = DriveItem("root-folder", "Corpus", FOLDER_MIME_TYPE, "drive")
    document = DriveItem(
        "document",
        "Plan",
        "application/vnd.google-apps.document",
        "drive",
        parents=("root-folder",),
    )
    gateway = SearchGateway(
        root,
        document,
        search_results=[],
        blobs={document.id: b"# Plan\n\nShip it."},
    )
    service = DocumentService(ScopedDrive(gateway, SHARED_DRIVE, "root-folder"))
    [listed] = (await service.list_folder()).items

    result = await service.read_document(listed.id)

    assert result.content == "# Plan\n\nShip it."
    assert gateway.exports == [(document.id, "text/markdown")]


@pytest.mark.anyio
async def test_reads_prefer_natural_sheet_boundaries() -> None:
    workbook = Workbook()
    workbook.active["A1"] = "ABC"
    workbook.create_sheet("Second")["A1"] = "DEFG"
    content = BytesIO()
    workbook.save(content)

    root = DriveItem("root-folder", "Corpus", FOLDER_MIME_TYPE, "drive")
    document = DriveItem(
        "document",
        "data.xlsx",
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "drive",
        parents=("root-folder",),
    )
    gateway = SearchGateway(
        root, document, search_results=[], blobs={document.id: content.getvalue()}
    )
    service = DocumentService(ScopedDrive(gateway, SHARED_DRIVE, "root-folder"))
    [listed] = (await service.list_folder()).items

    first = await service.read_document(listed.id, max_chars=6)
    second = await service.read_document(listed.id, cursor=first.next_cursor, max_chars=6)

    assert first.content == "ABC\n\n"
    assert first.location == 'sheet "Sheet", rows 1-1'
    assert second.content == "DEFG"
    assert second.location == 'sheet "Second", rows 1-1'


@pytest.mark.anyio
async def test_folder_listing_rejects_a_negative_cursor() -> None:
    root = DriveItem("root-folder", "Corpus", FOLDER_MIME_TYPE, "drive")
    document = DriveItem("document", "notes.txt", "text/plain", "drive", parents=("root-folder",))
    service = DocumentService(
        ScopedDrive(SearchGateway(root, document, search_results=[]), SHARED_DRIVE, "root-folder")
    )

    # "LTE" is base64 for "-1": as a slice bound it would page from the end.
    with pytest.raises(InvalidCursor):
        await service.list_folder(cursor="LTE")


@pytest.mark.anyio
async def test_read_rejects_an_invalid_cursor() -> None:
    root = DriveItem("root-folder", "Corpus", FOLDER_MIME_TYPE, "drive")
    document = DriveItem("document", "notes.txt", "text/plain", "drive", parents=("root-folder",))
    gateway = SearchGateway(root, document, search_results=[], blobs={document.id: b"notes"})
    service = DocumentService(ScopedDrive(gateway, SHARED_DRIVE, "root-folder"))
    [listed] = (await service.list_folder()).items

    with pytest.raises(InvalidCursor, match="Invalid read cursor"):
        await service.read_document(listed.id, cursor="A")


@pytest.mark.anyio
async def test_cached_content_is_not_returned_after_a_document_moves_outside() -> None:
    drive_root = DriveItem("drive", "Shared drive", FOLDER_MIME_TYPE, "drive")
    root = DriveItem("root-folder", "Corpus", FOLDER_MIME_TYPE, "drive", parents=("drive",))
    outside = DriveItem("outside", "Outside", FOLDER_MIME_TYPE, "drive", parents=("drive",))
    document = DriveItem(
        "document",
        "notes.txt",
        "text/plain",
        "drive",
        parents=("root-folder",),
        modified_time="2026-08-30T10:00:00Z",
    )
    gateway = SearchGateway(
        drive_root,
        root,
        outside,
        document,
        search_results=[],
        blobs={document.id: b"cached content"},
    )
    service = DocumentService(ScopedDrive(gateway, SHARED_DRIVE, "root-folder"))
    [listed] = (await service.list_folder()).items
    await service.read_document(listed.id)
    gateway.items[document.id] = replace(document, parents=(outside.id,))

    with pytest.raises(ScopeViolation):
        await service.read_document(listed.id)


@pytest.mark.anyio
async def test_search_enumerates_the_corpus_once_instead_of_walking_folder_by_folder() -> None:
    root = DriveItem("root-folder", "Corpus", FOLDER_MIME_TYPE, "drive")
    folders = [
        DriveItem(
            f"folder-{index}",
            f"Folder {index}",
            FOLDER_MIME_TYPE,
            "drive",
            parents=("root-folder",),
        )
        for index in range(20)
    ]
    gateway = SearchGateway(root, *folders, search_results=[])
    service = DocumentService(ScopedDrive(gateway, SHARED_DRIVE, "root-folder"))

    await service.search_documents("budget")

    assert gateway.folder_enumerations == 1
    assert gateway.child_listings == []


@pytest.mark.anyio
async def test_one_read_performs_one_ancestry_proof() -> None:
    root = DriveItem("root-folder", "Corpus", FOLDER_MIME_TYPE, "drive")
    plans = DriveItem("plans", "Plans", FOLDER_MIME_TYPE, "drive", parents=("root-folder",))
    document = DriveItem(
        "document",
        "notes.txt",
        "text/plain",
        "drive",
        parents=("plans",),
        modified_time="2026-08-30T10:00:00Z",
    )
    gateway = SearchGateway(
        root, plans, document, search_results=[document], blobs={document.id: b"notes"}
    )
    service = DocumentService(ScopedDrive(gateway, SHARED_DRIVE, "root-folder"))
    [found] = (await service.search_documents("notes")).items
    gateway.metadata_reads.clear()

    cold = await service.read_document(found.id)
    cold_reads = list(gateway.metadata_reads)
    gateway.metadata_reads.clear()
    warm = await service.read_document(found.id)

    assert cold.relative_path == warm.relative_path == "Plans/notes.txt"
    assert cold_reads == [document.id]
    # A cached parse still costs exactly one proof: the cache holds the bytes, never the
    # authorization decision.
    assert gateway.metadata_reads == [document.id]
    assert gateway.downloads == [document.id]


def text_document(name: str, body: bytes) -> tuple[DriveItem, DriveItem]:
    root = DriveItem("root-folder", "Corpus", FOLDER_MIME_TYPE, "drive")
    item = DriveItem("doc", name, "text/plain", "drive", parents=("root-folder",))
    return root, item


@pytest.mark.anyio
async def test_a_chunk_never_exceeds_the_byte_ceiling_even_in_cjk() -> None:
    """`max_chars` counts characters and the transport counts bytes. The Agent
    Engine SSE stream corrupts at roughly 100 KB per tool result, and CJK text
    is three bytes per character — so a 40,000-character chunk is 120 KB and
    would break the stream while satisfying every character limit.
    """
    root, item = text_document("notes.txt", "承認済".encode() * 20_000)
    gateway = SearchGateway(
        root, item, search_results=[], blobs={item.id: "承認済".encode() * 20_000}
    )
    service = DocumentService(ScopedDrive(gateway, SHARED_DRIVE, "root-folder"))

    result = await service.read_document(item.id, max_chars=40_000)

    assert len(result.content.encode("utf-8")) <= service.max_read_bytes
    assert result.next_cursor is not None


@pytest.mark.anyio
async def test_a_partial_read_says_that_it_is_partial() -> None:
    """A model not told it saw part of a document answers as though it saw all
    of it, and nothing in the content itself reveals the difference."""
    body = b"a" * 1_000
    root, item = text_document("notes.txt", body)
    gateway = SearchGateway(root, item, search_results=[], blobs={item.id: body})
    service = DocumentService(ScopedDrive(gateway, SHARED_DRIVE, "root-folder"))

    result = await service.read_document(item.id, max_chars=100)

    assert result.total_chars == 1_000
    assert result.note is not None
    assert "of 1,000" in result.note
    assert "next_cursor" in result.note


@pytest.mark.anyio
async def test_a_complete_read_carries_no_note() -> None:
    """The note is a warning, so it must not appear when nothing is missing."""
    body = b"short"
    root, item = text_document("notes.txt", body)
    gateway = SearchGateway(root, item, search_results=[], blobs={item.id: body})
    service = DocumentService(ScopedDrive(gateway, SHARED_DRIVE, "root-folder"))

    result = await service.read_document(item.id)

    assert result.note is None
    assert result.next_cursor is None
    assert result.total_chars == len("short")


@pytest.mark.anyio
async def test_the_parsed_cache_evicts_instead_of_growing_without_bound() -> None:
    """The dict this replaced never evicted: fine in a stdio process that exits
    with the session, unbounded in a server instance that lives for days."""
    root = DriveItem("root-folder", "Corpus", FOLDER_MIME_TYPE, "drive")
    items = [
        DriveItem(f"doc-{index}", f"{index}.txt", "text/plain", "drive", parents=("root-folder",))
        for index in range(3)
    ]
    gateway = SearchGateway(
        root, *items, search_results=[], blobs={item.id: b"body" for item in items}
    )
    service = DocumentService(ScopedDrive(gateway, SHARED_DRIVE, "root-folder"), cache_size=2)

    for item in items:
        await service.read_document(item.id)
    # The first document was evicted, so reading it again re-downloads.
    await service.read_document(items[0].id)

    assert gateway.downloads == ["doc-0", "doc-1", "doc-2", "doc-0"]


@pytest.mark.anyio
async def test_a_cached_document_is_not_downloaded_twice() -> None:
    body = b"body"
    root, item = text_document("notes.txt", body)
    gateway = SearchGateway(root, item, search_results=[], blobs={item.id: body})
    service = DocumentService(ScopedDrive(gateway, SHARED_DRIVE, "root-folder"))

    await service.read_document(item.id)
    await service.read_document(item.id)

    assert gateway.downloads == ["doc"]


@pytest.mark.anyio
async def test_an_oversized_blob_is_refused_before_it_is_downloaded() -> None:
    """Drive's own 10 MB cap applies to Workspace exports; an uploaded file can
    be any size, and pulling a gigabyte into an instance is its own outage."""
    root = DriveItem("root-folder", "Corpus", FOLDER_MIME_TYPE, "drive")
    huge = DriveItem(
        "doc", "huge.txt", "text/plain", "drive", parents=("root-folder",), size=40 * 1024 * 1024
    )
    gateway = SearchGateway(root, huge, search_results=[], blobs={huge.id: b"never read"})
    service = DocumentService(ScopedDrive(gateway, SHARED_DRIVE, "root-folder"))

    with pytest.raises(ExportTooLarge, match="40.0 MB"):
        await service.read_document(huge.id)

    assert gateway.downloads == []


@pytest.mark.anyio
async def test_a_scanned_pdf_is_reported_as_empty_rather_than_unsupported() -> None:
    """A scanned PDF parses perfectly and yields whitespace. Calling that
    "unsupported" sends the caller looking for a parser that already ran."""
    root = DriveItem("root-folder", "Corpus", FOLDER_MIME_TYPE, "drive")
    scan = DriveItem("doc", "scan.pdf", "application/pdf", "drive", parents=("root-folder",))
    gateway = SearchGateway(root, scan, search_results=[], blobs={scan.id: b"%PDF-"})
    service = DocumentService(
        ScopedDrive(gateway, SHARED_DRIVE, "root-folder"),
        extractors=_registry_returning(ExtractedDocument((TextBlock("page 1", "\n\n  \n"),))),
    )

    with pytest.raises(EmptyDocument, match="no OCR"):
        await service.read_document(scan.id)


@pytest.mark.anyio
async def test_an_empty_text_file_does_not_blame_a_scanner() -> None:
    """The two cases need different messages, or the PDF wording becomes noise
    that a reader learns to ignore."""
    body = b"   \n\n "
    root, item = text_document("empty.txt", body)
    gateway = SearchGateway(root, item, search_results=[], blobs={item.id: body})
    service = DocumentService(ScopedDrive(gateway, SHARED_DRIVE, "root-folder"))

    with pytest.raises(EmptyDocument) as raised:
        await service.read_document(item.id)

    assert "OCR" not in str(raised.value)


@pytest.mark.anyio
async def test_max_chars_is_capped_below_what_the_transport_can_carry() -> None:
    body = b"a" * 10
    root, item = text_document("notes.txt", body)
    gateway = SearchGateway(root, item, search_results=[], blobs={item.id: body})
    service = DocumentService(ScopedDrive(gateway, SHARED_DRIVE, "root-folder"))

    with pytest.raises(ValueError, match="between 1 and 40000"):
        await service.read_document(item.id, max_chars=50_000)


@pytest.mark.anyio
async def test_the_byte_ceiling_holds_when_block_boundaries_pull_the_cut_around() -> None:
    """The ceiling loop and the block-boundary preference move `end` in
    opposite directions; this pins that they still converge, and under the
    ceiling rather than just near it."""
    root = DriveItem("root-folder", "Corpus", FOLDER_MIME_TYPE, "drive")
    slides = DriveItem("doc", "deck.pdf", "application/pdf", "drive", parents=("root-folder",))
    blocks = tuple(TextBlock(f"page {index + 1}", "承認済" * 4_000) for index in range(10))
    gateway = SearchGateway(root, slides, search_results=[], blobs={slides.id: b"%PDF-"})
    service = DocumentService(
        ScopedDrive(gateway, SHARED_DRIVE, "root-folder"),
        extractors=_registry_returning(ExtractedDocument(blocks)),
    )

    cursor: str | None = None
    seen = 0
    for _ in range(50):
        result = await service.read_document(slides.id, cursor=cursor, max_chars=40_000)
        assert len(result.content.encode("utf-8")) <= service.max_read_bytes
        seen += len(result.content)
        cursor = result.next_cursor
        if cursor is None:
            break

    assert cursor is None, "paging did not terminate"
    assert seen == result.total_chars


@pytest.mark.anyio
async def test_keyword_search_asks_each_batch_for_one_more_than_the_limit() -> None:
    """One extra hit is all `truncated` needs; a full page of a thousand was paying
    for a count nobody sees."""
    root = DriveItem("root-folder", "Corpus", FOLDER_MIME_TYPE, "drive")
    gateway = SearchGateway(root, search_results=[])
    service = DocumentService(ScopedDrive(gateway, SHARED_DRIVE, "root-folder"))

    await service.search_documents("budget", limit=7)

    assert gateway.searched_limits == [8]


@pytest.mark.anyio
async def test_keyword_search_is_truncated_when_drive_had_more_to_give() -> None:
    root = DriveItem("root-folder", "Corpus", FOLDER_MIME_TYPE, "drive")
    hits = [
        DriveItem(
            f"hit-{index}", f"Hit {index}.md", "text/markdown", "drive", parents=("root-folder",)
        )
        for index in range(2)
    ]
    # Exactly `limit` hits came back, so the count alone would say "not truncated";
    # Drive's continuation token is what says otherwise.
    gateway = SearchGateway(root, *hits, search_results=hits, has_more=True)
    service = DocumentService(ScopedDrive(gateway, SHARED_DRIVE, "root-folder"))

    results = await service.search_documents("budget", limit=2)

    assert len(results.items) == 2
    assert results.truncated is True


@pytest.mark.anyio
async def test_keyword_search_is_not_truncated_when_every_match_was_returned() -> None:
    root = DriveItem("root-folder", "Corpus", FOLDER_MIME_TYPE, "drive")
    hit = DriveItem("hit", "Hit.md", "text/markdown", "drive", parents=("root-folder",))
    gateway = SearchGateway(root, hit, search_results=[hit])
    service = DocumentService(ScopedDrive(gateway, SHARED_DRIVE, "root-folder"))

    results = await service.search_documents("budget", limit=2)

    assert results.truncated is False


def _operations(caplog: pytest.LogCaptureFixture) -> list[logging.LogRecord]:
    return [
        record
        for record in caplog.records
        if record.name == "gdrive_scoped.audit" and hasattr(record, "event")
    ]


def _field(record: logging.LogRecord, name: str) -> object:
    """An `extra` field. Records carry them as attributes the type checker cannot see."""

    return getattr(record, name)


@pytest.mark.anyio
async def test_every_operation_leaves_one_audit_record_naming_the_caller(
    caplog: pytest.LogCaptureFixture,
) -> None:
    """Decisions say what the boundary did; operations say what the agent read.
    Both are needed to answer "what did it read, and for whom"."""
    root = DriveItem("root-folder", "Corpus", FOLDER_MIME_TYPE, "drive")
    report = DriveItem("report", "Budget.md", "text/markdown", "drive", parents=("root-folder",))
    gateway = SearchGateway(
        root, report, search_results=[report], blobs={"report": b"# Budget\n\nnumbers"}
    )
    service = DocumentService(ScopedDrive(gateway, SHARED_DRIVE, "root-folder"))

    token = audit_caller.set("alice@example.org")
    try:
        with caplog.at_level(logging.INFO, logger="gdrive_scoped.audit"):
            await service.search_documents("budget", limit=5)
            await service.list_folder()
            await service.get_metadata("report")
            await service.read_document("report")
    finally:
        audit_caller.reset(token)

    operations = _operations(caplog)
    assert [_field(record, "event") for record in operations] == [
        "search",
        "list_folder",
        "get_metadata",
        "read_document",
    ]
    assert {_field(record, "caller") for record in operations} == {"alice@example.org"}
    search, listing, metadata, read = operations
    assert (_field(search, "query"), _field(search, "file_ids"), _field(search, "truncated")) == (
        "budget",
        ["report"],
        False,
    )
    assert (_field(listing, "folder_id"), _field(listing, "file_ids")) == (
        "root-folder",
        ["report"],
    )
    assert _field(metadata, "file_id") == "report"
    assert (_field(read, "file_id"), _field(read, "chars"), _field(read, "partial")) == (
        "report",
        17,
        False,
    )


@pytest.mark.anyio
async def test_an_operation_without_a_known_caller_is_still_recorded(
    caplog: pytest.LogCaptureFixture,
) -> None:
    root = DriveItem("root-folder", "Corpus", FOLDER_MIME_TYPE, "drive")
    gateway = SearchGateway(root, search_results=[])
    service = DocumentService(ScopedDrive(gateway, SHARED_DRIVE, "root-folder"))

    with caplog.at_level(logging.INFO, logger="gdrive_scoped.audit"):
        await service.search_documents("budget")

    (record,) = _operations(caplog)
    assert _field(record, "caller") is None
