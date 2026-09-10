from __future__ import annotations

from io import BytesIO

import pytest
from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from pypdf import PdfWriter

from gdrive_scoped.extractors import ExtractorRegistry, UnsupportedContentType

XLSX_MIME_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
PPTX_MIME_TYPE = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
DOCX_MIME_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
PDF_MIME_TYPE = "application/pdf"


def test_workbook_extraction_reads_every_sheet() -> None:
    workbook = Workbook()
    first = workbook.active
    first.title = "Summary"
    first.append(["Metric", "Value"])
    first.append(["Revenue", 42])
    details = workbook.create_sheet("Details")
    details.append(["Region", "Owner"])
    details.append(["EMEA", "Alex"])
    content = BytesIO()
    workbook.save(content)

    plan = ExtractorRegistry().plan_for(XLSX_MIME_TYPE)
    extracted = plan.extractor(content.getvalue())

    assert [block.location for block in extracted.blocks] == [
        'sheet "Summary", rows 1-2',
        'sheet "Details", rows 1-2',
    ]
    assert "Metric\tValue\nRevenue\t42" in extracted.text
    assert "Region\tOwner\nEMEA\tAlex" in extracted.text


def test_presentation_extraction_preserves_slides_and_speaker_notes() -> None:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Roadmap"
    slide.placeholders[1].text = "Ship the PoC"
    notes = slide.notes_slide.notes_text_frame
    assert notes is not None
    notes.text = "Mention the August milestone."
    content = BytesIO()
    presentation.save(content)

    plan = ExtractorRegistry().plan_for(PPTX_MIME_TYPE)
    extracted = plan.extractor(content.getvalue())

    assert [block.location for block in extracted.blocks] == ["slide 1"]
    assert "Roadmap" in extracted.text
    assert "Ship the PoC" in extracted.text
    assert "Speaker notes:\nMention the August milestone." in extracted.text


def test_word_extraction_preserves_sections_and_tables() -> None:
    document = Document()
    document.add_heading("Overview", level=1)
    document.add_paragraph("This is the scoped Drive PoC.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Owner"
    table.cell(0, 1).text = "Status"
    table.cell(1, 0).text = "Alex"
    table.cell(1, 1).text = "Active"
    content = BytesIO()
    document.save(content)

    plan = ExtractorRegistry().plan_for(DOCX_MIME_TYPE)
    extracted = plan.extractor(content.getvalue())

    assert [block.location for block in extracted.blocks] == ["section: Overview", "table 1"]
    assert "Overview\nThis is the scoped Drive PoC." in extracted.text
    assert "Owner\tStatus\nAlex\tActive" in extracted.text


def test_pdf_extraction_preserves_page_locations() -> None:
    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    writer.add_blank_page(width=612, height=792)
    content = BytesIO()
    writer.write(content)

    plan = ExtractorRegistry().plan_for(PDF_MIME_TYPE)
    extracted = plan.extractor(content.getvalue())

    assert [block.location for block in extracted.blocks] == ["page 1", "page 2"]


@pytest.mark.parametrize(
    ("drive_mime_type", "export_mime_type"),
    [
        ("application/vnd.google-apps.document", "text/markdown"),
        ("application/vnd.google-apps.spreadsheet", XLSX_MIME_TYPE),
        ("application/vnd.google-apps.presentation", PPTX_MIME_TYPE),
    ],
)
def test_google_workspace_formats_choose_structure_preserving_exports(
    drive_mime_type: str, export_mime_type: str
) -> None:
    assert ExtractorRegistry().plan_for(drive_mime_type).export_mime_type == export_mime_type


def test_unsupported_content_remains_an_explicit_read_error() -> None:
    with pytest.raises(UnsupportedContentType, match="image/png"):
        ExtractorRegistry().plan_for("image/png")


@pytest.mark.parametrize(
    "mime_type",
    ["text/tab-separated-values", "text/xml", "application/xml"],
)
def test_types_the_census_found_now_pass_through_as_text(mime_type: str) -> None:
    """Added from a real census rather than guessed at. Their markup is the
    content, so passing it through is the honest thing to do."""
    plan = ExtractorRegistry().plan_for(mime_type)

    extracted = plan.extractor(b"<row><a>1</a></row>\tsecond")

    assert "second" in extracted.text


def test_html_is_reduced_to_its_text() -> None:
    """A page is mostly tags. Passing them through would spend most of a chunk
    on angle brackets and CSS rather than on anything a reader wants."""
    page = (
        b"<html><head><title>t</title><style>body{color:red}</style></head>"
        b"<body><h1>Quarterly report</h1><p>Revenue rose.</p>"
        b"<script>alert(1)</script><p>Costs fell.</p></body></html>"
    )

    extracted = ExtractorRegistry().plan_for("text/html").extractor(page)

    assert "Quarterly report" in extracted.text
    assert "Revenue rose." in extracted.text
    assert "Costs fell." in extracted.text
    assert "color:red" not in extracted.text
    assert "alert(1)" not in extracted.text
    assert "<" not in extracted.text


def test_html_keeps_block_boundaries_as_paragraph_breaks() -> None:
    """Otherwise adjacent paragraphs run into one another as a single word.

    One blank line between blocks, matching how the rest of the library joins
    them — and runs of empty tags collapse rather than accumulating.
    """
    page = b"<p>first</p><div></div><div></div><p>second</p>"

    extracted = ExtractorRegistry().plan_for("text/html").extractor(page)

    assert extracted.text == "first\n\nsecond"


@pytest.mark.parametrize(
    "mime_type",
    ["text/x-log", "text/yaml", "application/x-yaml", "application/yaml", "application/x-ndjson"],
)
def test_any_text_type_passes_through_even_when_nobody_listed_it(mime_type: str) -> None:
    """A corpus holds whatever people uploaded. A `text/*` type nobody thought of
    is legible as it stands, and refusing it sends a model looking for a parser
    that would only have decoded bytes."""
    registry = ExtractorRegistry()

    assert registry.supports(mime_type)
    extracted = registry.plan_for(mime_type).extractor(b"key: value\n")
    assert extracted.text == "key: value\n"


def test_the_text_prefix_rule_does_not_override_a_dedicated_extractor() -> None:
    page = b"<html><body><p>Hello</p><script>x()</script></body></html>"

    extracted = ExtractorRegistry().plan_for("text/html").extractor(page)

    assert extracted.text == "Hello"
