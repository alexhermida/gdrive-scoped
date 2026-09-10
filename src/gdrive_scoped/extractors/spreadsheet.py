"""XLSX extraction with every worksheet and stable row locations."""

from __future__ import annotations

from io import BytesIO
from typing import Any

from openpyxl import load_workbook

from gdrive_scoped.extractors.models import ExtractedDocument, TextBlock

ROWS_PER_BLOCK = 100


def extract_workbook(content: bytes) -> ExtractedDocument:
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    blocks: list[TextBlock] = []
    try:
        for worksheet in workbook.worksheets:
            rows: list[str] = []
            start_row = 1
            for row_number, values in enumerate(worksheet.iter_rows(values_only=True), start=1):
                if not rows:
                    start_row = row_number
                rows.append(_render_row(values))
                if len(rows) == ROWS_PER_BLOCK:
                    blocks.append(_sheet_block(worksheet.title, start_row, row_number, rows))
                    rows = []
            if rows:
                end_row = start_row + len(rows) - 1
                blocks.append(_sheet_block(worksheet.title, start_row, end_row, rows))
    finally:
        workbook.close()
    return ExtractedDocument(tuple(blocks))


def _render_row(values: tuple[Any, ...]) -> str:
    cells = ["" if value is None else str(value) for value in values]
    while cells and not cells[-1]:
        cells.pop()
    return "\t".join(cells)


def _sheet_block(title: str, start: int, end: int, rows: list[str]) -> TextBlock:
    return TextBlock(
        location=f'sheet "{title}", rows {start}-{end}',
        text="\n".join(rows),
    )
