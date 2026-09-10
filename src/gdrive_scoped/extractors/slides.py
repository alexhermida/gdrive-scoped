"""PPTX extraction that keeps slide and speaker-note boundaries."""

from __future__ import annotations

from io import BytesIO
from typing import Any, cast

from pptx import Presentation

from gdrive_scoped.extractors.models import ExtractedDocument, TextBlock


def extract_presentation(content: bytes) -> ExtractedDocument:
    presentation = Presentation(BytesIO(content))
    blocks: list[TextBlock] = []
    for slide_number, slide in enumerate(presentation.slides, start=1):
        lines: list[str] = []
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            text = cast("Any", shape).text_frame.text.strip()
            if text:
                lines.append(text)

        notes_frame = slide.notes_slide.notes_text_frame if slide.has_notes_slide else None
        notes = notes_frame.text.strip() if notes_frame is not None else ""
        if notes:
            lines.extend(("Speaker notes:", notes))

        blocks.append(TextBlock(location=f"slide {slide_number}", text="\n".join(lines)))
    return ExtractedDocument(tuple(blocks))
